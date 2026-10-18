"""Write a PPTX from slides.json. Text is already computed; this file only lays it out.

Optional slide fields: stats [{value, label}], image (one picture), images [{src, caption}].
Picture paths are relative to the folder that holds slides.json.
"""

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

src, dest = sys.argv[1], sys.argv[2]
base = Path(src).resolve().parent
slides = json.load(open(src, encoding="utf-8"))["slides"]
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
ink = RGBColor(0x1C, 0x19, 0x15)
teal = RGBColor(0x0F, 0x6E, 0x6E)
muted = RGBColor(0x5C, 0x56, 0x4C)
paper = RGBColor(0xF4, 0xF1, 0xEA)
card = RGBColor(0xFF, 0xFD, 0xF8)
line = RGBColor(0xE4, 0xDD, 0xD0)
night = RGBColor(0x0F, 0x2F, 0x2F)
mint = RGBColor(0x8F, 0xD3, 0xC8)
cream = RGBColor(0xE7, 0xE2, 0xD6)


def text(page, left, top, width, height, value, size, color, font="Calibri", bold=False):
    box = page.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.text = value
    para.font.size = Pt(size)
    para.font.color.rgb = color
    para.font.name = font
    para.font.bold = bold
    return frame


def picture(page, path, left, top, max_w, max_h):
    pic = page.shapes.add_picture(str(base / path), Inches(left), Inches(top))
    scale = min(Inches(max_w) / pic.width, Inches(max_h) / pic.height)
    pic.width = int(pic.width * scale)
    pic.height = int(pic.height * scale)
    pic.left = int(Inches(left) + (Inches(max_w) - pic.width) / 2)
    return pic.height / 914400


for index, slide in enumerate(slides, start=1):
    cover = index == 1
    page = prs.slides.add_slide(blank)
    fill = page.background.fill
    fill.solid()
    fill.fore_color.rgb = night if cover else paper
    text(page, 0.7, 0.36, 12, 0.32, slide["kicker"].upper(), 12, mint if cover else teal, bold=True)
    text(page, 0.7, 0.66, 12, 0.9, slide["title"], 48 if cover else 32, paper if cover else ink, font="Georgia")
    y = 1.75 if cover else 1.5
    stats = slide.get("stats") or []
    if stats:
        gap = 0.2
        width = (11.93 - gap * (len(stats) - 1)) / len(stats)
        for i, item in enumerate(stats):
            left = 0.7 + i * (width + gap)
            shape = page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(width), Inches(1.45))
            shape.adjustments[0] = 0.12
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(0x1E, 0x45, 0x45) if cover else card
            shape.line.color.rgb = RGBColor(0x3A, 0x60, 0x60) if cover else line
            text(page, left + 0.12, y + 0.08, width - 0.24, 0.6, item["value"], 26, mint if cover else teal, font="Georgia", bold=True)
            text(page, left + 0.12, y + 0.66, width - 0.24, 0.75, item["label"], 11, cream if cover else muted)
        y += 1.65
    if slide.get("image"):
        y += picture(page, slide["image"], 0.7, y, 11.93, 4.75) + 0.12
    images = slide.get("images") or []
    if images:
        gap = 0.2
        width = (11.93 - gap * (len(images) - 1)) / len(images)
        heights = [picture(page, item["src"], 0.7 + i * (width + gap), y, width, 3.6) for i, item in enumerate(images)]
        tallest = max(heights)
        for i, item in enumerate(images):
            text(page, 0.7 + i * (width + gap), y + tallest + 0.08, width, 0.5, item["caption"], 12, muted)
        y += tallest + 0.75
    bullets = slide["bullets"]
    dense = len(bullets) > 6 or (stats and len(bullets) > 4)
    body = page.shapes.add_textbox(Inches(0.85), Inches(y), Inches(11.6), Inches(max(0.6, 6.95 - y)))
    frame = body.text_frame
    frame.word_wrap = True
    for i, entry in enumerate(bullets):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = f"•  {entry}"
        para.level = 0
        para.font.size = Pt(14 if dense else 17)
        para.font.color.rgb = cream if cover else ink
        para.font.name = "Calibri"
        para.space_after = Pt(6)
    text(
        page,
        0.7,
        7.05,
        12,
        0.3,
        f"Research prototype — not a medical device, not medical advice.    {index} / {len(slides)}",
        11,
        cream if cover else muted,
    )

prs.save(dest)
print(f"Wrote {dest} ({len(slides)} slides)")
